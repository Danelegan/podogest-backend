import hmac
import re
from zoneinfo import ZoneInfo

from django.conf import settings
from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_GET, require_http_methods
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response

from .models import Appointment, ClinicalRecord
from .serializers import AppointmentSerializer, ClinicalRecordSerializer
from .services import read_appointment_token, send_appointment_confirmation, send_due_reminders
from .throttles import ReservaBurstThrottle, ReservaSustainedThrottle


# Encabezados del respaldo y su ancho máximo de columna en Excel.
BACKUP_COLUMNS = [
    'Fecha', 'Hora', 'Estado', 'Nombre', 'RUT', 'Email', 'Teléfono',
    'Antecedentes', 'Síntomas', 'Diagnóstico', 'Tratamiento', 'Observaciones',
]
BACKUP_MAX_WIDTH = 50

# El frontend guarda el podograma dentro de Observaciones como
# "<texto>|||DRAWING:data:image/png;base64,...".
DRAWING_RE = re.compile(r'\|\|\|DRAWING:\S*')


def _clean_observaciones(text):
    """Quita la imagen en base64 del podograma y deja solo el texto de la podóloga."""
    if not text or '|||DRAWING:' not in text:
        return text
    text = DRAWING_RE.sub('', text).strip()
    note = '[Podograma guardado en sistema]'
    return f'{text}\n{note}' if text else note


def _excel_text(value):
    # openpyxl rechaza caracteres de control que podrían venir del formulario público.
    return ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value


class AppointmentViewSet(viewsets.ModelViewSet):
    queryset = Appointment.objects.all()
    serializer_class = AppointmentSerializer

    def get_permissions(self):
        # Agendar y consultar horarios ocupados es público; el resto exige autenticación.
        if self.action in ['create', 'horarios_ocupados']:
            return [AllowAny()]
        return [IsAuthenticated()]

    def get_throttles(self):
        # Crear reservas es público: se añade un límite estricto contra spam/DoS
        # además del límite anónimo general. Los usuarios con JWT no se ven afectados.
        if self.action == 'create':
            return [*super().get_throttles(), ReservaBurstThrottle(), ReservaSustainedThrottle()]
        return super().get_throttles()

    @action(detail=False, methods=['get'], url_path='horarios-ocupados')
    def horarios_ocupados(self, request):
        # Solo fecha y hora, sin datos personales, para bloquearlas en el calendario.
        citas = (
            self.get_queryset()
            .exclude(status=Appointment.Status.CANCELLED)
            .values('appointment_date', 'appointment_time')
        )
        return Response(list(citas))

    def perform_create(self, serializer):
        appointment = serializer.save()
        send_appointment_confirmation(appointment)

    def perform_update(self, serializer):
        # Cubre PUT y PATCH desde el panel: al cancelar, misma limpieza que el enlace del correo.
        with transaction.atomic():
            appointment = serializer.save()
            if appointment.status == Appointment.Status.CANCELLED:
                appointment.delete_empty_clinical_record()

    @action(detail=False, methods=['get'], url_path='exportar-respaldo')
    def exportar_respaldo(self, request):
        """Descarga todas las citas con su ficha en un Excel (respaldo manual)."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Respaldo'
        sheet.append(BACKUP_COLUMNS)

        for appointment in self.get_queryset().select_related('clinical_record'):
            # Las citas sin ficha dejan esas columnas en blanco.
            record = getattr(appointment, 'clinical_record', None)
            clinical = [getattr(record, field) if record else None for field in ClinicalRecord.CLINICAL_FIELDS]
            clinical[-1] = _clean_observaciones(clinical[-1])
            sheet.append([_excel_text(value) for value in [
                appointment.appointment_date,
                appointment.appointment_time,
                appointment.get_status_display(),
                appointment.patient_name,
                appointment.rut,
                appointment.email,
                appointment.phone,
                *clinical,
            ]])

        for cell in sheet['A'][1:]:
            cell.number_format = 'DD-MM-YYYY'
        for cell in sheet['B'][1:]:
            cell.number_format = 'HH:MM'
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # Los datos vienen del formulario público: un texto como "=HYPERLINK(...)"
                # se guarda como texto y no como fórmula.
                if cell.data_type == 'f':
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)

        header_fill = PatternFill('solid', fgColor='0284C7')
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.freeze_panes = 'A2'

        for column in sheet.columns:
            lengths = [
                max(len(line) for line in str(cell.value).splitlines() or [''])
                for cell in column if cell.value is not None
            ]
            width = min(max(lengths, default=0) + 2, BACKUP_MAX_WIDTH)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)

        today = timezone.localtime(timezone.now(), ZoneInfo(settings.CLINIC_TIME_ZONE)).date()
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = (
            f'attachment; filename="respaldo_pasos_saludables_{today:%Y%m%d}.xlsx"'
        )
        workbook.save(response)
        return response

    @action(detail=False, methods=['post'], url_path='atencion-espontanea')
    def atencion_espontanea(self, request):
        """Registra a un paciente que llega sin reserva: crea la cita ahora mismo y su ficha."""
        # Hora local de la clínica; con segundos para no chocar con los bloques
        # reservados (que van en horas exactas) ni con otro espontáneo del mismo minuto.
        now = timezone.localtime(timezone.now(), ZoneInfo(settings.CLINIC_TIME_ZONE))
        serializer = self.get_serializer(data={
            'patient_name': request.data.get('patient_name'),
            'rut': request.data.get('rut'),
            'phone': request.data.get('phone'),
            'email': request.data.get('email', ''),
            'service_type': request.data.get('service_type') or Appointment.ServiceType.GENERAL,
            'appointment_date': now.date(),
            'appointment_time': now.time().replace(microsecond=0),
            # El paciente ya está presente: no requiere confirmación.
            'status': Appointment.Status.CONFIRMED,
        })
        serializer.is_valid(raise_exception=True)

        # Sin correo de confirmación: la atención es inmediata. Tampoco recibe
        # recordatorio, porque el cron solo toma citas que aún no empiezan.
        with transaction.atomic():
            appointment = serializer.save()
            record = ClinicalRecord.objects.create(appointment=appointment)

        return Response({
            'appointment': serializer.data,
            'clinical_record': ClinicalRecordSerializer(record).data,
        }, status=status.HTTP_201_CREATED)


class ClinicalRecordViewSet(viewsets.ModelViewSet):
    """Fichas podológicas: solo accesibles con JWT (uso exclusivo de la podóloga)."""

    queryset = ClinicalRecord.objects.select_related('appointment')
    serializer_class = ClinicalRecordSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = super().get_queryset()
        # Permite buscar la ficha de una cita: /api/clinical-records/?appointment=<id>
        appointment_id = self.request.query_params.get('appointment')
        if appointment_id:
            queryset = queryset.filter(appointment_id=appointment_id)
        return queryset


# --- Enlaces del correo de recordatorio -------------------------------------
# Vistas HTML simples (fuera de DRF): el paciente no tiene JWT, el token firmado
# de la URL es lo que lo autoriza.

def _appointment_from_token(token):
    pk = read_appointment_token(token)
    if pk is None:
        return None
    return Appointment.objects.filter(pk=pk).first()


def _render_result(request, title, message, status=200):
    return render(request, 'citas/accion_cita.html', {
        'title': title,
        'message': message,
    }, status=status)


def _invalid_link(request):
    return _render_result(
        request,
        'Enlace no válido',
        'Este enlace no es válido o la cita ya no existe. Si necesitas ayuda, responde el correo que te enviamos.',
        status=404,
    )


def _past_appointment(request):
    return _render_result(
        request,
        'Cita finalizada',
        'Esta cita ya pasó, por lo que no es posible modificarla.',
        status=410,
    )


@require_GET
def confirmar_cita(request, token):
    appointment = _appointment_from_token(token)
    if appointment is None:
        return _invalid_link(request)
    if appointment.status == Appointment.Status.CANCELLED:
        return _render_result(
            request,
            'Cita cancelada',
            'Esta cita fue cancelada anteriormente. Si deseas asistir, agenda una nueva hora en nuestro sitio web.',
        )
    if appointment.starts_at <= timezone.now():
        return _past_appointment(request)

    if appointment.status != Appointment.Status.CONFIRMED:
        appointment.status = Appointment.Status.CONFIRMED
        appointment.save(update_fields=['status'])

    return _render_result(
        request,
        '¡Cita confirmada!',
        f'Gracias, {appointment.patient_name}. Te esperamos el {appointment.appointment_date:%d-%m-%Y} '
        f'a las {appointment.appointment_time:%H:%M}.',
    )


# El token firmado de la URL reemplaza la protección CSRF (no hay sesión que proteger).
@csrf_exempt
@require_http_methods(['GET', 'POST'])
def cancelar_cita(request, token):
    appointment = _appointment_from_token(token)
    if appointment is None:
        return _invalid_link(request)
    if appointment.status == Appointment.Status.CANCELLED:
        return _render_result(request, 'Cita cancelada', 'Esta cita ya se encontraba cancelada.')
    if appointment.starts_at <= timezone.now():
        return _past_appointment(request)

    # GET solo pide confirmación: algunos filtros de correo abren los enlaces
    # automáticamente y no deben poder cancelar la cita por sí solos.
    if request.method == 'GET':
        return render(request, 'citas/accion_cita.html', {
            'title': '¿Cancelar tu cita?',
            'message': (
                f'{appointment.patient_name}, vas a cancelar tu cita del '
                f'{appointment.appointment_date:%d-%m-%Y} a las {appointment.appointment_time:%H:%M}.'
            ),
            'confirm_cancel': True,
        })

    with transaction.atomic():
        appointment.status = Appointment.Status.CANCELLED
        appointment.save(update_fields=['status'])
        # El paciente no asistirá: una ficha sin datos clínicos solo sería basura.
        appointment.delete_empty_clinical_record()
    return _render_result(
        request,
        'Cita cancelada',
        'Tu cita fue cancelada correctamente. Puedes agendar una nueva hora cuando lo necesites.',
    )


# --- Endpoint para el cron externo (cron-job.org) ----------------------------

def _cron_token(request):
    auth = request.headers.get('Authorization', '')
    if auth.startswith('Bearer '):
        return auth.removeprefix('Bearer ').strip()
    return request.headers.get('X-Cron-Token') or request.GET.get('token', '')


# Lo llama un servicio externo sin sesión: se protege con CRON_SECRET, no con CSRF.
@csrf_exempt
@require_http_methods(['GET', 'POST'])
def cron_send_reminders(request):
    secret = settings.CRON_SECRET
    # Sin CRON_SECRET configurado el endpoint queda deshabilitado.
    # compare_digest evita filtrar el token midiendo el tiempo de respuesta.
    if not secret or not hmac.compare_digest(_cron_token(request).encode(), secret.encode()):
        return JsonResponse({'detail': 'No autorizado.'}, status=401)

    sent, due = send_due_reminders()
    return JsonResponse({'enviados': sent, 'pendientes': due})
