import importlib
import io
from datetime import date, time, timedelta
from unittest.mock import patch
from zoneinfo import ZoneInfo

from django.conf import settings
from django.contrib.auth.models import User
from django.core.cache import cache
from django.db import OperationalError
from django.test import TestCase, override_settings
from django.urls import clear_url_caches, reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework.throttling import AnonRateThrottle

from api.models import Appointment, ClinicalRecord
from api.services import make_appointment_token
from api.views import _clean_observaciones
import core.urls


def reserva(hour):
    return {
        'patient_name': 'Paciente Prueba',
        'rut': '11111111-1',
        'phone': '+56911111111',
        'service_type': 'general',
        'appointment_date': date(2030, 1, 7).isoformat(),
        'appointment_time': time(hour, 0).isoformat(),
    }


class ReservaThrottleTests(TestCase):
    def setUp(self):
        cache.clear()
        self.client = APIClient()

    def test_reservas_anonimas_limitadas_a_5_por_minuto(self):
        for hour in range(9, 14):
            response = self.client.post('/api/appointments/', reserva(hour), format='json')
            self.assertEqual(response.status_code, 201)

        response = self.client.post('/api/appointments/', reserva(14), format='json')
        self.assertEqual(response.status_code, 429)

    def test_horarios_ocupados_no_usa_el_limite_de_reservas(self):
        for _ in range(10):
            response = self.client.get('/api/appointments/horarios-ocupados/')
            self.assertEqual(response.status_code, 200)

    def test_usuario_autenticado_no_tiene_limite_de_reservas(self):
        user = User.objects.create_user('podologa', password='clave-segura-123')
        self.client.force_authenticate(user)
        for hour in range(9, 16):
            response = self.client.post('/api/appointments/', reserva(hour), format='json')
            self.assertEqual(response.status_code, 201)


class CorsTests(TestCase):
    def preflight(self, origin):
        return self.client.options(
            '/api/appointments/',
            HTTP_ORIGIN=origin,
            HTTP_ACCESS_CONTROL_REQUEST_METHOD='POST',
        )

    def test_origen_permitido(self):
        response = self.preflight('https://podogest-frontend.vercel.app')
        self.assertEqual(
            response.headers.get('Access-Control-Allow-Origin'),
            'https://podogest-frontend.vercel.app',
        )

    def test_origen_no_permitido(self):
        response = self.preflight('https://sitio-malicioso.com')
        self.assertNotIn('Access-Control-Allow-Origin', response.headers)


@override_settings(CRON_SECRET='secreto-de-prueba')
class RecordatorioTests(TestCase):
    def crear_cita(self, en_horas, **extra):
        from api.models import Appointment
        inicio = timezone.localtime(
            timezone.now() + timedelta(hours=en_horas),
            ZoneInfo(settings.CLINIC_TIME_ZONE),
        )
        datos = {
            'patient_name': 'Paciente Prueba',
            'rut': '11111111-1',
            'phone': '+56911111111',
            'email': 'paciente@example.com',
            'service_type': 'general',
            'appointment_date': inicio.date(),
            'appointment_time': inicio.time().replace(second=0, microsecond=0),
        }
        datos.update(extra)
        return Appointment.objects.create(**datos)

    @patch('api.services.resend.Emails.send', return_value={'id': 'test'})
    def test_envia_solo_citas_dentro_de_3_horas_y_una_vez(self, send):
        cita = self.crear_cita(2.5)
        self.crear_cita(5)  # demasiado lejos
        self.crear_cita(1, email='')  # sin correo
        self.crear_cita(2, status='cancelada')

        url = reverse('cron_send_reminders')
        response = self.client.get(url, HTTP_AUTHORIZATION='Bearer secreto-de-prueba')
        self.assertEqual(response.json(), {'enviados': 1, 'pendientes': 1})
        response = self.client.get(url, {'token': 'secreto-de-prueba'})
        self.assertEqual(response.json(), {'enviados': 0, 'pendientes': 0})

        self.assertEqual(send.call_count, 1)
        params = send.call_args.args[0]
        self.assertEqual(params['to'], ['paciente@example.com'])
        self.assertEqual(params['from'], 'PasosSaludables <contacto@pasossaludables.cl>')
        self.assertIn('Hola <strong>Paciente Prueba</strong>', params['html'])
        self.assertIn(reverse('confirmar_cita', args=[make_appointment_token(cita)]), params['html'])
        cita.refresh_from_db()
        self.assertIsNotNone(cita.reminder_sent_at)

    def test_confirmar(self):
        cita = self.crear_cita(2)
        response = self.client.get(reverse('confirmar_cita', args=[make_appointment_token(cita)]))
        self.assertEqual(response.status_code, 200)
        cita.refresh_from_db()
        self.assertEqual(cita.status, 'confirmada')

    def test_cancelar_requiere_post_y_libera_el_horario(self):
        cita = self.crear_cita(2)
        url = reverse('cancelar_cita', args=[make_appointment_token(cita)])

        self.client.get(url)
        cita.refresh_from_db()
        self.assertEqual(cita.status, 'pendiente')

        self.client.post(url)
        cita.refresh_from_db()
        self.assertEqual(cita.status, 'cancelada')
        # El horario cancelado se puede volver a reservar.
        self.crear_cita(2, appointment_date=cita.appointment_date, appointment_time=cita.appointment_time)

    def test_cancelar_elimina_la_ficha_vacia(self):
        cita = self.crear_cita(2)
        ClinicalRecord.objects.create(appointment=cita, observaciones='   ')
        self.client.post(reverse('cancelar_cita', args=[make_appointment_token(cita)]))
        cita.refresh_from_db()
        self.assertEqual(cita.status, 'cancelada')
        self.assertFalse(ClinicalRecord.objects.filter(appointment=cita).exists())

    def test_cancelar_conserva_la_ficha_con_datos(self):
        cita = self.crear_cita(2)
        ClinicalRecord.objects.create(appointment=cita, antecedentes_medicos='Diabetes tipo 2')
        self.client.post(reverse('cancelar_cita', args=[make_appointment_token(cita)]))
        cita.refresh_from_db()
        self.assertEqual(cita.status, 'cancelada')
        self.assertTrue(ClinicalRecord.objects.filter(appointment=cita).exists())

    def test_token_invalido(self):
        cita = self.crear_cita(2)
        response = self.client.get(reverse('confirmar_cita', args=[f'{cita.pk}:falso']))
        self.assertEqual(response.status_code, 404)

    @patch('api.services.resend.Emails.send')
    def test_cron_rechaza_token_incorrecto_o_ausente(self, send):
        self.crear_cita(2)
        url = reverse('cron_send_reminders')
        self.assertEqual(self.client.get(url).status_code, 401)
        self.assertEqual(self.client.get(url, {'token': 'otro'}).status_code, 401)
        send.assert_not_called()

    @override_settings(CRON_SECRET='')
    def test_cron_deshabilitado_sin_secreto(self):
        response = self.client.get(reverse('cron_send_reminders'), {'token': ''})
        self.assertEqual(response.status_code, 401)


class AtencionEspontaneaTests(TestCase):
    url = '/api/appointments/atencion-espontanea/'
    datos = {
        'patient_name': 'Paciente Espontáneo',
        'rut': '22222222-2',
        'phone': '+56922222222',
        'email': 'espontaneo@example.com',
    }

    def setUp(self):
        self.client = APIClient()

    def test_requiere_autenticacion(self):
        response = self.client.post(self.url, self.datos, format='json')
        self.assertEqual(response.status_code, 401)

    @patch('api.services.resend.Emails.send')
    def test_crea_cita_ahora_y_su_ficha(self, send):
        self.client.force_authenticate(User.objects.create_user('podologa'))
        response = self.client.post(self.url, self.datos, format='json')
        self.assertEqual(response.status_code, 201)

        cita = Appointment.objects.get(pk=response.data['appointment']['id'])
        self.assertEqual(cita.status, 'confirmada')
        self.assertEqual(cita.service_type, 'general')
        self.assertLess(abs((timezone.now() - cita.starts_at).total_seconds()), 5)
        self.assertEqual(response.data['clinical_record']['id'], cita.clinical_record.pk)
        send.assert_not_called()

    def test_datos_invalidos_no_crean_nada(self):
        self.client.force_authenticate(User.objects.create_user('podologa'))
        response = self.client.post(self.url, {**self.datos, 'email': 'no-es-correo'}, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Appointment.objects.exists())


class EliminarCitaTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.cita = Appointment.objects.create(
            patient_name='Paciente Prueba', rut='11111111-1', phone='+56911111111',
            service_type='general', appointment_date=date(2030, 1, 7), appointment_time=time(10, 0),
        )
        ClinicalRecord.objects.create(appointment=self.cita, diagnostico='Onicomicosis')
        self.url = f'/api/appointments/{self.cita.pk}/'

    def test_requiere_autenticacion(self):
        self.assertEqual(self.client.delete(self.url).status_code, 401)
        self.assertTrue(Appointment.objects.exists())

    def test_elimina_la_cita_y_su_ficha(self):
        self.client.force_authenticate(User.objects.create_user('podologa'))
        self.assertEqual(self.client.delete(self.url).status_code, 204)
        self.assertFalse(Appointment.objects.exists())
        self.assertFalse(ClinicalRecord.objects.exists())


class CancelarDesdePanelTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_user('podologa'))
        self.cita = Appointment.objects.create(
            patient_name='Paciente Prueba', rut='11111111-1', phone='+56911111111',
            service_type='general', appointment_date=date(2030, 1, 7), appointment_time=time(10, 0),
        )
        self.url = f'/api/appointments/{self.cita.pk}/'

    def test_cancelar_elimina_la_ficha_vacia(self):
        ClinicalRecord.objects.create(appointment=self.cita)
        response = self.client.patch(self.url, {'status': 'cancelada'}, format='json')
        self.assertEqual(response.status_code, 200)
        self.cita.refresh_from_db()
        self.assertEqual(self.cita.status, 'cancelada')
        self.assertFalse(ClinicalRecord.objects.exists())

    def test_cancelar_conserva_la_ficha_con_datos(self):
        ClinicalRecord.objects.create(appointment=self.cita, sintomas='Dolor al caminar')
        self.client.patch(self.url, {'status': 'cancelada'}, format='json')
        self.assertTrue(ClinicalRecord.objects.exists())

    def test_otros_cambios_no_tocan_la_ficha(self):
        ClinicalRecord.objects.create(appointment=self.cita)
        self.client.patch(self.url, {'status': 'confirmada'}, format='json')
        self.assertTrue(ClinicalRecord.objects.exists())


class FichaUnicaPorCitaTests(TestCase):
    url = '/api/clinical-records/'

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_user('podologa'))
        self.cita = Appointment.objects.create(
            patient_name='Paciente Prueba', rut='11111111-1', phone='+56911111111',
            service_type='general', appointment_date=date(2030, 1, 7), appointment_time=time(10, 0),
        )

    def test_segunda_ficha_para_la_misma_cita_es_rechazada(self):
        primera = self.client.post(self.url, {'appointment': self.cita.pk, 'sintomas': 'Dolor'}, format='json')
        self.assertEqual(primera.status_code, 201)

        segunda = self.client.post(self.url, {'appointment': self.cita.pk, 'sintomas': 'Otro'}, format='json')
        self.assertEqual(segunda.status_code, 400)
        self.assertIn('appointment', segunda.data)
        self.assertEqual(ClinicalRecord.objects.filter(appointment=self.cita).count(), 1)

    def test_filtra_fichas_por_cita(self):
        otra_cita = Appointment.objects.create(
            patient_name='Otro Paciente', rut='22222222-2', phone='+56922222222',
            service_type='general', appointment_date=date(2030, 1, 7), appointment_time=time(11, 0),
        )
        ficha = ClinicalRecord.objects.create(appointment=self.cita, diagnostico='Onicomicosis')
        ClinicalRecord.objects.create(appointment=otra_cita, diagnostico='Verruga plantar')

        response = self.client.get(self.url, {'appointment': self.cita.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]['id'], ficha.pk)
        self.assertEqual(response.data[0]['appointment'], self.cita.pk)

        # Sin el parámetro se listan todas las fichas.
        self.assertEqual(len(self.client.get(self.url).data), 2)

    def test_filtro_por_cita_no_numerico_devuelve_400(self):
        response = self.client.get(self.url, {'appointment': 'abc'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data, {'appointment': 'Debe ser el ID numérico de la cita.'})


class ExportarRespaldoTests(TestCase):
    url = '/api/appointments/exportar-respaldo/'

    def setUp(self):
        self.client = APIClient()
        datos = {'rut': '11111111-1', 'phone': '912345678', 'service_type': 'general',
                 'appointment_date': date(2030, 1, 7)}
        con_ficha = Appointment.objects.create(
            patient_name='Ana Pérez', email='ana@example.com', appointment_time=time(10, 0), **datos,
        )
        ClinicalRecord.objects.create(appointment=con_ficha, diagnostico='Onicomicosis')
        Appointment.objects.create(
            patient_name='=HYPERLINK("http://malo")', appointment_time=time(11, 0),
            status='cancelada', **datos,
        )

    def test_requiere_autenticacion(self):
        self.assertEqual(self.client.get(self.url).status_code, 401)

    def test_descarga_citas_con_su_ficha_en_excel(self):
        ClinicalRecord.objects.filter(diagnostico='Onicomicosis').update(
            observaciones='Uña engrosada|||DRAWING:data:image/png;base64,iVBORw0KGgoAAAANSUhEUg==',
        )
        self.client.force_authenticate(User.objects.create_user('podologa'))
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertRegex(
            response['Content-Disposition'],
            r'attachment; filename="respaldo_pasos_saludables_\d{8}\.xlsx"',
        )

        sheet = load_workbook(io.BytesIO(response.content)).active
        filas = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(filas[0][:3], ['Fecha', 'Hora', 'Estado'])
        self.assertEqual(sheet['A1'].font.color.rgb, '00FFFFFF')
        self.assertEqual(filas[1][0].date(), date(2030, 1, 7))
        self.assertEqual(filas[1][1], time(10, 0))
        self.assertEqual(filas[1][2:], [
            'Pendiente', 'Ana Pérez', '11111111-1', 'ana@example.com', '912345678',
            None, None, 'Onicomicosis', None,
            'Uña engrosada\n[Podograma guardado en sistema]',
        ])
        # Sin ficha: columnas clínicas vacías; el texto que parece fórmula queda como texto.
        self.assertEqual(filas[2][2:4], ['Cancelada', '=HYPERLINK("http://malo")'])
        self.assertEqual(sheet['D3'].data_type, 's')
        self.assertEqual(filas[2][7:], [None] * 5)

    def test_ningun_texto_se_exporta_como_formula(self):
        formula = '=HYPERLINK("http://malo","clic")'
        cita = Appointment.objects.create(
            patient_name=formula, rut='=1+1', phone='=1+1', service_type='general',
            appointment_date=date(2030, 1, 8), appointment_time=time(9, 0),
        )
        ClinicalRecord.objects.create(
            appointment=cita, **{field: formula for field in ClinicalRecord.CLINICAL_FIELDS},
        )
        self.client.force_authenticate(User.objects.create_user('podologa'))
        sheet = load_workbook(io.BytesIO(self.client.get(self.url).content)).active

        formulas = [cell.coordinate for row in sheet.iter_rows() for cell in row if cell.data_type == 'f']
        self.assertEqual(formulas, [])
        fila = next(row for row in sheet.iter_rows(min_row=2) if row[4].value == '=1+1')
        self.assertEqual([cell.value for cell in fila[7:]], [formula] * 5)

    def test_observaciones_solo_con_podograma(self):
        self.assertEqual(
            _clean_observaciones('|||DRAWING:data:image/png;base64,AAAA'),
            '[Podograma guardado en sistema]',
        )
        self.assertEqual(_clean_observaciones('Sin dibujo'), 'Sin dibujo')

    def test_expone_el_nombre_del_archivo_al_frontend(self):
        self.client.force_authenticate(User.objects.create_user('podologa'))
        response = self.client.get(self.url, HTTP_ORIGIN='https://podogest-frontend.vercel.app')
        self.assertIn('Content-Disposition', response['Access-Control-Expose-Headers'])


class HealthCheckTests(TestCase):
    url = '/api/health/'

    def test_responde_ok_sin_autenticacion(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {'status': 'ok'})

    def test_responde_503_si_la_base_de_datos_falla(self):
        with patch('api.views.connection.cursor', side_effect=OperationalError('db caída')),                 self.assertLogs('api.views', level='ERROR'):
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 503)
        # No filtra el error de la base de datos.
        self.assertEqual(response.json(), {'status': 'error'})

    def test_no_le_aplica_el_limite_anonimo(self):
        cache.clear()
        for _ in range(105):
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)


class ReservaPublicaTests(TestCase):
    url = '/api/appointments/'
    # Exactamente lo que envía BookingModal.jsx del frontend.
    formulario = {**reserva(10), 'email': 'paciente@example.com'}

    def setUp(self):
        cache.clear()
        self.client = APIClient()

    @patch('api.services.resend.Emails.send')
    def test_acepta_el_formulario_del_sitio(self, send):
        response = self.client.post(self.url, self.formulario, format='json')
        self.assertEqual(response.status_code, 201)
        cita = Appointment.objects.get()
        self.assertEqual(cita.status, 'pendiente')
        self.assertEqual(cita.email, 'paciente@example.com')
        send.assert_called_once()

    @patch('api.services.resend.Emails.send')
    def test_ignora_campos_internos(self, send):
        response = self.client.post(self.url, {
            **self.formulario,
            'status': 'confirmada',
            'reminder_sent_at': '2030-01-01T00:00:00Z',
        }, format='json')
        self.assertEqual(response.status_code, 201)
        self.assertNotIn('status', response.data)
        cita = Appointment.objects.get()
        self.assertEqual(cita.status, 'pendiente')
        self.assertIsNone(cita.reminder_sent_at)

    @patch('api.services.resend.Emails.send')
    def test_el_panel_autenticado_puede_fijar_el_estado(self, send):
        self.client.force_authenticate(User.objects.create_user('podologa'))
        response = self.client.post(self.url, {**self.formulario, 'status': 'confirmada'}, format='json')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(Appointment.objects.get().status, 'confirmada')


class LoginThrottleTests(TestCase):
    url = '/api/token/'
    credenciales = {'username': 'podologa', 'password': 'incorrecta'}

    def setUp(self):
        cache.clear()
        self.client = APIClient()
        self.ahora = 1_000_000.0
        reloj = patch('rest_framework.throttling.SimpleRateThrottle.timer', lambda _: self.ahora)
        reloj.start()
        self.addCleanup(reloj.stop)

    def intentar(self):
        return self.client.post(self.url, self.credenciales, format='json').status_code

    def test_login_limitado_a_5_por_minuto(self):
        for _ in range(5):
            self.assertEqual(self.intentar(), 401)
        self.assertEqual(self.intentar(), 429)
        self.ahora += 61
        self.assertEqual(self.intentar(), 401)

    def test_login_limitado_a_20_por_hora(self):
        for _ in range(4):
            for _ in range(5):
                self.assertEqual(self.intentar(), 401)
            self.ahora += 61
        self.assertEqual(self.intentar(), 429)
        self.ahora += 3600
        self.assertEqual(self.intentar(), 401)

    def test_login_correcto_sigue_funcionando(self):
        User.objects.create_user('podologa', password='clave-segura-123')
        response = self.client.post(
            self.url, {'username': 'podologa', 'password': 'clave-segura-123'}, format='json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('access', response.data)

    def test_separado_del_cupo_anonimo(self):
        with patch.dict(AnonRateThrottle.THROTTLE_RATES, {'anon': '1/min'}):
            self.assertEqual(self.client.get('/api/appointments/horarios-ocupados/').status_code, 200)
            self.assertEqual(self.client.get('/api/appointments/horarios-ocupados/').status_code, 429)
            # El cupo anónimo agotado no bloquea el login...
            self.assertEqual(self.intentar(), 401)
        # ...y el login tampoco gasta el cupo anónimo.
        cache.clear()
        for _ in range(5):
            self.intentar()
        with patch.dict(AnonRateThrottle.THROTTLE_RATES, {'anon': '1/min'}):
            self.assertEqual(self.client.get('/api/appointments/horarios-ocupados/').status_code, 200)


class AdminSoloEnDesarrolloTests(TestCase):
    def rutas_con_debug(self, debug):
        with override_settings(DEBUG=debug):
            rutas = [str(p.pattern) for p in importlib.reload(core.urls).urlpatterns]
        importlib.reload(core.urls)
        clear_url_caches()
        return rutas

    def test_admin_no_existe_en_produccion(self):
        self.assertEqual(self.client.get('/admin/').status_code, 404)
        self.assertNotIn('admin/', self.rutas_con_debug(False))

    def test_admin_existe_en_desarrollo(self):
        self.assertIn('admin/', self.rutas_con_debug(True))


class PermisosTests(TestCase):
    """Con IsAuthenticated por defecto, solo lo público a propósito queda abierto."""

    def setUp(self):
        cache.clear()
        self.client = APIClient()
        self.cita = Appointment.objects.create(
            patient_name='Paciente Prueba', rut='11111111-1', phone='+56911111111',
            service_type='general', appointment_date=date(2030, 1, 7), appointment_time=time(10, 0),
        )
        self.ficha = ClinicalRecord.objects.create(appointment=self.cita)

    # --- Públicos a propósito ---

    def test_crear_reserva_es_publico(self):
        response = self.client.post('/api/appointments/', reserva(11), format='json')
        self.assertEqual(response.status_code, 201)

    def test_horarios_ocupados_es_publico(self):
        self.assertEqual(self.client.get('/api/appointments/horarios-ocupados/').status_code, 200)

    def test_confirmar_y_cancelar_solo_piden_el_token_firmado(self):
        token = make_appointment_token(self.cita)
        self.assertEqual(self.client.get(reverse('cancelar_cita', args=[token])).status_code, 200)
        self.assertEqual(self.client.get(reverse('confirmar_cita', args=[token])).status_code, 200)
        self.assertEqual(self.client.get(reverse('confirmar_cita', args=['malo'])).status_code, 404)

    @override_settings(CRON_SECRET='secreto-de-prueba')
    @patch('api.views.send_due_reminders', return_value=(0, 0))
    def test_cron_solo_pide_su_secreto(self, _):
        url = reverse('cron_send_reminders')
        self.assertEqual(self.client.get(url).status_code, 401)
        response = self.client.get(url, HTTP_AUTHORIZATION='Bearer secreto-de-prueba')
        self.assertEqual(response.status_code, 200)

    def test_health_es_publico(self):
        self.assertEqual(self.client.get('/api/health/').status_code, 200)

    def test_login_y_refresh_son_publicos(self):
        User.objects.create_user('podologa', password='clave-segura-123')
        login = self.client.post('/api/token/', {'username': 'podologa', 'password': 'clave-segura-123'},
                                 format='json')
        self.assertEqual(login.status_code, 200)
        refresh = self.client.post('/api/token/refresh/', {'refresh': login.data['refresh']}, format='json')
        self.assertEqual(refresh.status_code, 200)

    # --- Todo lo demás exige JWT ---

    def test_el_resto_exige_jwt(self):
        cita, ficha = f'/api/appointments/{self.cita.pk}/', f'/api/clinical-records/{self.ficha.pk}/'
        peticiones = [
            ('get', '/api/'),
            ('get', '/api/appointments/'),
            ('get', cita),
            ('patch', cita),
            ('put', cita),
            ('delete', cita),
            ('get', '/api/appointments/exportar-respaldo/'),
            ('post', '/api/appointments/atencion-espontanea/'),
            ('get', '/api/clinical-records/'),
            ('post', '/api/clinical-records/'),
            ('get', ficha),
            ('patch', ficha),
            ('put', ficha),
            ('delete', ficha),
        ]
        for metodo, url in peticiones:
            with self.subTest(metodo=metodo, url=url):
                response = getattr(self.client, metodo)(url, {}, format='json')
                self.assertEqual(response.status_code, 401)
        self.assertTrue(Appointment.objects.filter(pk=self.cita.pk).exists())
        self.assertTrue(ClinicalRecord.objects.filter(pk=self.ficha.pk).exists())

    def test_con_jwt_real_el_panel_funciona(self):
        User.objects.create_user('podologa', password='clave-segura-123')
        login = self.client.post('/api/token/', {'username': 'podologa', 'password': 'clave-segura-123'},
                                 format='json')
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {login.data["access"]}')
        self.assertEqual(self.client.get('/api/appointments/').status_code, 200)
        self.assertEqual(self.client.get('/api/clinical-records/').status_code, 200)
        self.assertEqual(self.client.get('/api/').status_code, 200)
